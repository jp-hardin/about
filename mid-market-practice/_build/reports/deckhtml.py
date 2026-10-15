# Reads a PowerPoint deck for the site: the theme colors, every shape's position, fill and outline, its text with
# inherited styles, links, tables and pictures, and each chart redrawn as SVG from its data (with a data table).
# webdeck.py arranges what this reads into the site's report sections. Positions and sizes are in points.
import hashlib, html, io, os, re
from decimal import Decimal, ROUND_HALF_UP
from datetime import date, timedelta
from lxml import etree
from PIL import Image
from pptx.opc.constants import RELATIONSHIP_TYPE as RT

NS = {
    'a': 'http://schemas.openxmlformats.org/drawingml/2006/main',
    'p': 'http://schemas.openxmlformats.org/presentationml/2006/main',
    'r': 'http://schemas.openxmlformats.org/officeDocument/2006/relationships',
    'c': 'http://schemas.openxmlformats.org/drawingml/2006/chart',
}
A, P, R, C = ('{%s}' % NS[k] for k in 'aprc')
EMU_PT = 12700
esc = html.escape


def X(el, path):
    return etree._Element.xpath(el, path, namespaces=NS) if el is not None else []


def X1(el, path):
    r = X(el, path)
    return r[0] if r else None


def n(v, d=2):
    """Compact number for CSS."""
    s = f'{v:.{d}f}'
    if '.' in s:
        s = s.rstrip('0').rstrip('.')
    return '0' if s in ('-0', '') else s


FONT_MAP = {  # deck typeface -> (CSS family, weight); the files are in assets/fonts/briefs/
    'calibri': ('Carlito', None), 'cambria': ('Caladea', None),
    'quicksand': ('Quicksand', 400), 'quicksand light': ('Quicksand', 300), 'quicksand medium': ('Quicksand', 500),
    'quicksand semibold': ('Quicksand', 600), 'quicksand bold': ('Quicksand', 700),
    'newsreader': ('Newsreader', 400), 'newsreader medium': ('Newsreader', 500), 'newsreader semibold': ('Newsreader', 600),
    'newsreader bold': ('Newsreader', 700), 'cinzel': ('Cinzel', 400), 'cinzel semibold': ('Cinzel', 600),
}
FALLBACK = {'Carlito': 'Calibri, Carlito, Arial, sans-serif', 'Caladea': 'Cambria, Caladea, Georgia, serif',
            'Quicksand': 'Quicksand, Arial, sans-serif', 'Newsreader': 'Newsreader, Georgia, serif', 'Cinzel': 'Cinzel, Georgia, serif'}


# ─── Colors ───────────────────────────────────────────────────────────────────
def _hex(h):
    h = h.lstrip('#')
    return [int(h[i:i + 2], 16) / 255 for i in (0, 2, 4)]


def _rgb_to_hsl(r, g, b):
    mx, mn = max(r, g, b), min(r, g, b)
    l = (mx + mn) / 2
    if mx == mn:
        return 0.0, 0.0, l
    d = mx - mn
    s = d / (2 - mx - mn) if l > 0.5 else d / (mx + mn)
    if mx == r:
        h = (g - b) / d + (6 if g < b else 0)
    elif mx == g:
        h = (b - r) / d + 2
    else:
        h = (r - g) / d + 4
    return h / 6, s, l


def _hsl_to_rgb(h, s, l):
    if s == 0:
        return l, l, l

    def f(p, q, t):
        t %= 1
        if t < 1 / 6:
            return p + (q - p) * 6 * t
        if t < 1 / 2:
            return q
        if t < 2 / 3:
            return p + (q - p) * (2 / 3 - t) * 6
        return p
    q = l * (1 + s) if l < 0.5 else l + s - l * s
    p = 2 * l - q
    return f(p, q, h + 1 / 3), f(p, q, h), f(p, q, h - 1 / 3)


class Theme:
    def __init__(self, master):
        theme = etree.fromstring(master.part.part_related_by(RT.THEME).blob)
        self.colors = {}
        for el in X(theme, './/a:clrScheme/*'):
            name = etree.QName(el).localname
            v = X1(el, './a:srgbClr/@val') or X1(el, './a:sysClr/@lastClr') or '000000'
            self.colors[name] = v
        cm = X1(master._element, './p:clrMap')
        self.map = dict(cm.attrib) if cm is not None else {}
        self.major = X1(theme, './/a:majorFont/a:latin/@typeface') or 'Calibri'
        self.minor = X1(theme, './/a:minorFont/a:latin/@typeface') or 'Calibri'

    def scheme(self, name):
        name = self.map.get(name, name)
        return self.colors.get(name, '000000')

    def color(self, el, ph_color=None):
        """CSS color for the first color child of el (solidFill, fontRef, ...). None for none."""
        if el is None:
            return None
        c = X1(el, './a:srgbClr|./a:schemeClr|./a:prstClr|./a:sysClr|./a:scrgbClr')
        if c is None:
            return None
        tag = etree.QName(c).localname
        if tag == 'srgbClr':
            base = c.get('val')
        elif tag == 'schemeClr':
            base = ph_color if c.get('val') == 'phClr' and ph_color else self.scheme(c.get('val'))
        elif tag == 'sysClr':
            base = c.get('lastClr', '000000')
        elif tag == 'prstClr':
            base = {'black': '000000', 'white': 'FFFFFF'}.get(c.get('val'), '000000')
        else:
            base = '000000'
        r, g, b = _hex(base)
        alpha = 1.0
        for m in c:
            k, v = etree.QName(m).localname, int(m.get('val', '0')) / 100000
            if k in ('lumMod', 'lumOff', 'tint', 'shade'):
                h, s, l = _rgb_to_hsl(r, g, b)
                if k == 'lumMod':
                    l *= v
                elif k == 'lumOff':
                    l += v
                elif k == 'tint':
                    l = l * v + (1 - v)
                elif k == 'shade':
                    l *= v
                r, g, b = _hsl_to_rgb(h, s, min(1, max(0, l)))
            elif k == 'alpha':
                alpha = v
        hx = '#%02x%02x%02x' % tuple(round(x * 255) for x in (r, g, b))
        return hx if alpha >= 0.999 else f'rgba({round(r * 255)},{round(g * 255)},{round(b * 255)},{n(alpha, 3)})'

    def fill(self, sppr, style=None):
        """CSS background for a shape's fill, or None."""
        if sppr is not None:
            if X1(sppr, './a:noFill') is not None:
                return None
            sf = X1(sppr, './a:solidFill')
            if sf is not None:
                return self.color(sf)
            gf = X1(sppr, './a:gradFill')
            if gf is not None:
                stops = ', '.join(f'{self.color(gs)} {n(int(gs.get("pos", 0)) / 1000)}%' for gs in X(gf, './a:gsLst/a:gs'))
                ang = int(X1(gf, './a:lin/@ang') or 0) / 60000
                return f'linear-gradient({n(ang + 90)}deg, {stops})'
        if style is not None:
            fr = X1(style, './a:fillRef')
            if fr is not None and fr.get('idx', '0') != '0':
                return self.color(fr)
        return None

    def line(self, sppr, style=None):
        """(width pt, CSS color, dash) for a shape's outline, or None."""
        ln = X1(sppr, './a:ln') if sppr is not None else None
        if ln is not None and X1(ln, './a:noFill') is not None:
            return None
        color = self.color(X1(ln, './a:solidFill')) if ln is not None else None
        if color is None and style is not None:
            lr = X1(style, './a:lnRef')
            if lr is not None and lr.get('idx', '0') != '0' and (ln is None or X1(ln, './a:solidFill') is None):
                color = self.color(lr)
        if color is None:
            return None
        w = int(ln.get('w', 9525)) / EMU_PT if ln is not None else 0.75
        dash = X1(ln, './a:prstDash/@val') if ln is not None else None
        return w, color, ('dashed' if dash and 'dash' in dash.lower() else 'dotted' if dash and 'dot' in dash.lower() else 'solid')


# ─── Number formats (charts) ─────────────────────────────────────────────────
def fmt(v, code):
    """Formats a value with an Excel number format (the common cases used in the decks)."""
    if v is None:
        return ''
    secs = (code or 'General').split(';')
    sec = secs[0]
    neg = ''
    if v < 0:
        if len(secs) > 1:
            sec, v = secs[1], -v
        else:
            neg, v = '-', -v
    m = re.match(r'^((?:"[^"]*"|\\.|[$€£+−-])*)(.*?)((?:"[^"]*"|\\.|[%xX ])*)$', sec)
    pre, core, post = (m.group(1), m.group(2), m.group(3)) if m else ('', sec, '')
    pct = '%' in re.sub(r'"[^"]*"', '', post)
    unq = lambda t: re.sub(r'"([^"]*)"', r'\1', t).replace('\\', '')
    pre, post = unq(pre), unq(post)
    scale = len(core) - len(core.rstrip(','))
    core = core.rstrip(',')
    dec = len(core.split('.')[1].rstrip('#').replace(',', '')) if '.' in core else 0
    if core in ('General', ''):
        dec = 0 if float(v).is_integer() else min(2, len(repr(float(v)).split('.')[1]))
    x = v * 100 if pct else v / 1000 ** scale
    x = float(Decimal(f'{x:.15g}').quantize(Decimal(1).scaleb(-dec), ROUND_HALF_UP))
    num = f'{x:,.{dec}f}' if ',' in core or abs(x) >= 10000 else f'{x:.{dec}f}'
    return neg + pre + num + post


def excel_date(serial, code):
    d = date(1899, 12, 30) + timedelta(days=int(serial))
    code = re.sub(r'"[^"]*"|\\.', '', code.split(';')[0]).lower()
    if 'mmm' in code:
        return d.strftime('%b %Y' if 'yyyy' in code else "%b '%y" if 'yy' in code else '%b')
    if 'yy' in code and 'm' not in code:
        return d.strftime('%Y')
    return d.strftime('%m/%Y')


# ─── Renderer ────────────────────────────────────────────────────────────────
class Box:
    """One rendered element with its position on the page (points)."""
    def __init__(self, x, y, w, h, html, cls='', container=None, chrome=False, deco=False, text=''):
        self.x, self.y, self.w, self.h = x, y, w, h
        self.html, self.cls, self.container, self.chrome, self.deco, self.text = html, cls, container, chrome, deco, text
        self.children = []


class Deck:
    def __init__(self, prs, media_dir, media_url, link_map=None):
        self.prs = prs
        self.W = prs.slide_width / EMU_PT
        self.H = prs.slide_height / EMU_PT
        self.master = prs.slide_masters[0]
        self.theme = Theme(self.master)
        self.media_dir, self.media_url = media_dir, media_url
        self.link_map = link_map or (lambda u: u)
        self.media = {}
        self.fonts_used = set()
        dts = X1(prs.part._element, './p:defaultTextStyle')
        self.default_style = dts
        os.makedirs(media_dir, exist_ok=True)

    # -- text style inheritance --------------------------------------------
    def _ph_chain(self, shape_el, slide_part):
        """lstStyle elements that a placeholder inherits from: layout placeholder, master placeholder, master text style."""
        ph = X1(shape_el, './p:nvSpPr/p:nvPr/p:ph')
        chain = []
        if ph is None:
            tx = X1(self.master._element, './p:txStyles/p:otherStyle')
            return [tx, self.default_style], None
        idx, typ = ph.get('idx'), ph.get('type', 'body')
        layout = slide_part.slide_layout if hasattr(slide_part, 'slide_layout') else None
        lay_sp = mas_sp = None
        if layout is not None:
            for sp in X(layout._element, './/p:sp'):
                lph = X1(sp, './p:nvSpPr/p:nvPr/p:ph')
                if lph is None:
                    continue
                if (idx is not None and lph.get('idx') == idx) or (idx is None and lph.get('type', 'body') == typ):
                    lay_sp = sp
                    break
        for sp in X(self.master._element, './/p:sp'):
            mph = X1(sp, './p:nvSpPr/p:nvPr/p:ph')
            if mph is not None and mph.get('type', 'body') == ('title' if typ in ('title', 'ctrTitle') else typ if typ in ('body', 'dt', 'ftr', 'sldNum') else 'body'):
                mas_sp = sp
                break
        for sp in (lay_sp, mas_sp):
            if sp is not None:
                chain.append(X1(sp, './p:txBody/a:lstStyle'))
        key = 'titleStyle' if typ in ('title', 'ctrTitle') else 'bodyStyle' if typ in ('body', 'subTitle', 'obj') else 'otherStyle'
        chain.append(X1(self.master._element, f'./p:txStyles/p:{key}'))
        chain.append(self.default_style)
        return chain, (lay_sp, mas_sp)

    def _lvl(self, chain, lvl):
        out = []
        for st in chain:
            if st is not None:
                e = X1(st, f'./a:lvl{lvl}pPr')
                if e is not None:
                    out.append(e)
        return out

    def _pattr(self, ppr, lvls, name, default=None):
        if ppr is not None and ppr.get(name) is not None:
            return ppr.get(name)
        for e in lvls:
            if e.get(name) is not None:
                return e.get(name)
        return default

    def _pchild(self, ppr, lvls, path):
        for e in ([ppr] if ppr is not None else []) + lvls:
            c = X1(e, path)
            if c is not None:
                return c
        return None

    def _rattr(self, rpr, lvls, name, default=None):
        if rpr is not None and rpr.get(name) is not None:
            return rpr.get(name)
        for e in lvls:
            d = X1(e, './a:defRPr')
            if d is not None and d.get(name) is not None:
                return d.get(name)
        return default

    def _rchild(self, rpr, lvls, path):
        if rpr is not None:
            c = X1(rpr, path)
            if c is not None:
                return c
        for e in lvls:
            d = X1(e, './a:defRPr')
            c = X1(d, path) if d is not None else None
            if c is not None:
                return c
        return None

    def font(self, face, bold):
        face = face or self.theme.minor
        if face.startswith('+mj'):
            face = self.theme.major
        elif face.startswith('+mn'):
            face = self.theme.minor
        fam, w = FONT_MAP.get(face.lower(), (face, None))
        weight = 700 if bold else (w or 400)
        if bold and w and w >= 500:
            weight = 700
        self.fonts_used.add((fam, weight))
        return FALLBACK.get(fam, f'"{fam}", sans-serif'), weight

    # -- text --------------------------------------------------------------
    def text_html(self, txbody, chain, rels, slide_no, scale=1.0, font_ref_color=None, title=False):
        paras, plain, meta = [], [], []
        for p in X(txbody, './a:p'):
            ppr = X1(p, './a:pPr')
            lvl = int(ppr.get('lvl', '0')) + 1 if ppr is not None else 1
            lvls = self._lvl(chain, lvl)
            st = []
            algn = self._pattr(ppr, lvls, 'algn', 'l')
            if algn != 'l':
                st.append('text-align:' + {'ctr': 'center', 'r': 'right', 'just': 'justify', 'dist': 'justify'}.get(algn, 'left'))
            marl = int(self._pattr(ppr, lvls, 'marL', '0')) / EMU_PT
            ind = int(self._pattr(ppr, lvls, 'indent', '0')) / EMU_PT
            if marl:
                st.append(f'padding-left:calc(var(--u)*{n(marl)})')
            if ind:
                st.append(f'text-indent:calc(var(--u)*{n(ind)})')
            for side, path in (('top', './a:spcBef'), ('bottom', './a:spcAft')):
                sp = self._pchild(ppr, lvls, path)
                if sp is not None:
                    pts = X1(sp, './a:spcPts/@val')
                    pct = X1(sp, './a:spcPct/@val')
                    if pts and int(pts):
                        st.append(f'margin-{side}:calc(var(--u)*{n(int(pts) / 100)})')
                    elif pct and int(pct):
                        st.append(f'margin-{side}:{n(int(pct) / 100000)}em')
            runs, sizes, texts, sem = [], [], [], []
            all_bold = all_italic = True
            colors, caps_seen = [], []
            for el in p:
                tag = etree.QName(el).localname
                if tag == 'br':
                    runs.append('<br>')
                    sem.append('<br>')
                    continue
                if tag not in ('r', 'fld'):
                    continue
                t = X1(el, './a:t')
                t = t.text if t is not None and t.text else ''
                if tag == 'fld' and el.get('type') == 'slidenum':
                    t = str(slide_no)
                if not t:
                    continue
                rpr = X1(el, './a:rPr')
                sz = int(self._rattr(rpr, lvls, 'sz', '1800')) / 100 * scale
                b = self._rattr(rpr, lvls, 'b', '0') in ('1', 'true')
                it = self._rattr(rpr, lvls, 'i', '0') in ('1', 'true')
                u = self._rattr(rpr, lvls, 'u', 'none')
                cap = self._rattr(rpr, lvls, 'cap', 'none')
                spc = self._rattr(rpr, lvls, 'spc')
                base = self._rattr(rpr, lvls, 'baseline')
                face = X1(rpr, './a:latin/@typeface') if rpr is not None else None
                if face is None:
                    lat = self._rchild(None, lvls, './a:latin')
                    face = lat.get('typeface') if lat is not None else ('+mj-lt' if title else '+mn-lt')
                fam, weight = self.font(face, b)
                color = self.theme.color(self._rchild(rpr, lvls, './a:solidFill')) or font_ref_color or self.theme.color(etree.fromstring(f'<x xmlns:a="{NS["a"]}"><a:schemeClr val="tx1"/></x>'))
                rs = [f'font-size:calc(var(--u)*{n(sz)})', f'font-weight:{weight}', f'font-family:{fam}', f'color:{color}']
                if it:
                    rs.append('font-style:italic')
                if u and u != 'none':
                    rs.append('text-decoration:underline')
                if cap == 'all':
                    rs.append('text-transform:uppercase')
                elif cap == 'small':
                    rs.append('font-variant:small-caps')
                if spc and int(spc):
                    rs.append(f'letter-spacing:calc(var(--u)*{n(int(spc) / 100)})')
                if base and int(base) > 0:
                    rs.append('vertical-align:super;font-size:0.7em')
                sizes.append(sz)
                texts.append(t)
                all_bold = all_bold and b
                all_italic = all_italic and it
                colors.append(color)
                if t.strip():
                    caps_seen.append(cap == 'all')
                inner = esc(t).replace('\x0b', '<br>')
                si = inner
                if b and t.strip():
                    si = f'<strong>{si}</strong>'
                if it and t.strip():
                    si = f'<em>{si}</em>'
                sem_link = None
                link = X1(rpr, './a:hlinkClick') if rpr is not None else None
                if link is not None:
                    rid = link.get(R + 'id')
                    url = rels[rid].target_ref if rid and rid in rels and rels[rid].is_external else None
                    action = link.get('action', '')
                    if url:
                        url = self.link_map(url)
                        ext = url.startswith('http')
                        tgt = ' target="_blank" rel="noopener"' if ext else ''
                        runs.append(f'<a href="{esc(url)}" style="{";".join(rs)}"{tgt}>{inner}</a>')
                        sem.append(f'<a href="{esc(url)}"{tgt}>{si}</a>')
                        continue
                    if 'hlinksldjump' in action and rid in rels:
                        target = rels[rid].target_part
                        num = next((i for i, s in enumerate(self.prs.slides, 1) if s.part is target), None)
                        if num:
                            runs.append(f'<a href="#p{num:02d}" style="{";".join(rs)}">{inner}</a>')
                            sem.append(f'<a href="#p{num:02d}">{si}</a>')
                            continue
                runs.append(f'<span style="{";".join(rs)}">{inner}</span>')
                sem.append(si)
            text = ''.join(texts)
            ln = self._pchild(ppr, lvls, './a:lnSpc')
            fs = max(sizes) if sizes else int(self._rattr(X1(p, './a:endParaRPr'), lvls, 'sz', '1800')) / 100 * scale
            lh = 1.2
            if ln is not None:
                pts, pct = X1(ln, './a:spcPts/@val'), X1(ln, './a:spcPct/@val')
                if pts:
                    lh = int(pts) / 100 / fs if fs else 1.2
                elif pct:
                    lh = int(pct) / 100000 * 1.2
            st.append(f'line-height:{n(lh, 3)}')
            semh = re.sub(r'</strong>(\s*)<strong>', r'\1', ''.join(sem))
            semh = re.sub(r'</em>(\s*)<em>', r'\1', semh)
            semh = re.sub(r'^(<br>)+|(<br>)+$', '', semh.strip()).strip()
            is_bullet = self._pchild(ppr, lvls, './a:buChar') is not None and X1(ppr, './a:buNone') is None if ppr is not None else self._pchild(ppr, lvls, './a:buChar') is not None
            meta.append(dict(sem=semh, text=text, size=fs, bold=all_bold and bool(texts), italic=all_italic and bool(texts),
                             color=colors[0] if colors else None, bullet=is_bullet and bool(text.strip()), align=algn, empty=not text.strip(),
                             caps=bool(caps_seen) and all(caps_seen)))
            if not text.strip():
                paras.append(f'<p class="e" style="{";".join(st)};font-size:calc(var(--u)*{n(fs)})">&nbsp;</p>')
                continue
            bu = self._pchild(ppr, lvls, './a:buChar')
            if bu is not None and X1(ppr, './a:buNone') is None:
                buc = self.theme.color(self._pchild(ppr, lvls, './a:buClr')) or ''
                runs.insert(0, f'<span class="bu" style="{"color:" + buc + ";" if buc else ""}display:inline-block;text-indent:0;width:calc(var(--u)*{n(abs(ind) if ind else fs)})">{esc(bu.get("char", "•"))}</span>')
            tag = 'p'
            if title:
                tag = 'h2'
            elif all_bold and fs >= 11 and len(text) < 110:
                tag = 'h3'
            paras.append(f'<{tag} style="{";".join(st)}">{"".join(runs)}</{tag}>')
            plain.append(text)
        self.last_meta = meta
        return '\n'.join(paras), ' '.join(plain)

    # -- pictures ----------------------------------------------------------
    def picture(self, blip_el, part, w, h, crop=None, alpha=None):
        rid = blip_el.get(R + 'embed')
        img_part = part.related_part(rid)
        blob = img_part.blob
        key = hashlib.sha1(blob + repr(crop).encode() + f'{round(w)}x{round(h)}'.encode()).hexdigest()[:12]
        if key in self.media:
            return self.media[key]
        ext = os.path.splitext(img_part.partname)[1].lower()
        if ext in ('.svg', '.emf', '.wmf'):
            name = f'{key}{ext}'
            open(os.path.join(self.media_dir, name), 'wb').write(blob)
        else:
            im = Image.open(io.BytesIO(blob))
            im.load()
            if crop:
                l, t, r, b = crop
                W0, H0 = im.size
                im = im.crop((round(W0 * l), round(H0 * t), round(W0 * (1 - r)), round(H0 * (1 - b))))
            has_alpha = im.mode in ('RGBA', 'LA', 'P') and 'transparency' in im.info or im.mode in ('RGBA', 'LA')
            im = im.convert('RGBA' if has_alpha else 'RGB')
            target = max(64, round(w * 2.4))   # two device pixels per CSS pixel at the widest layout
            if im.width > target:
                im = im.resize((target, max(1, round(im.height * target / im.width))), Image.LANCZOS)
            name = f'{key}.webp'
            im.save(os.path.join(self.media_dir, name), 'WEBP', quality=82, method=6)
        url = f'{self.media_url}/{name}'
        self.media[key] = url
        return url

    # -- shapes ------------------------------------------------------------
    def shapes(self, spTree, part, slide_no, xf=None, chrome=False, out=None, rels=None):
        """Walks a shape tree and returns Boxes in z-order. xf maps child coordinates to page points."""
        if out is None:
            out = []
        if xf is None:
            xf = lambda x, y: (x / EMU_PT, y / EMU_PT)
            sx = sy = 1 / EMU_PT
        else:
            sx, sy = xf.sx, xf.sy
        rels_ = rels if rels is not None else part.rels
        mark = None

        def tag_links(mark):
            if mark:
                for b in out[mark[0]:]:
                    if not getattr(b, 'href', None):
                        b.href = mark[1]
        for el in spTree:
            tag_links(mark)
            mark = None
            tag = etree.QName(el).localname
            if tag in ('sp', 'pic'):
                # a link set on the whole shape (a button), not on its text
                hl = X1(el, './*/p:cNvPr/a:hlinkClick')
                rid = hl.get(R + 'id') if hl is not None else None
                if rid and rid in rels_ and rels_[rid].is_external:
                    mark = (len(out), self.link_map(rels_[rid].target_ref))
            if tag == 'grpSp':
                g = X1(el, './p:grpSpPr/a:xfrm')
                off, ext = X1(g, './a:off'), X1(g, './a:ext')
                coff, cext = X1(g, './a:chOff'), X1(g, './a:chExt')
                ox, oy = int(off.get('x')), int(off.get('y'))
                ex, ey = int(ext.get('cx')), int(ext.get('cy'))
                cx0, cy0 = int(coff.get('x')), int(coff.get('y'))
                cex, cey = int(cext.get('cx')) or 1, int(cext.get('cy')) or 1
                parent = xf

                def child(x, y, parent=parent, ox=ox, oy=oy, ex=ex, ey=ey, cx0=cx0, cy0=cy0, cex=cex, cey=cey):
                    return parent(ox + (x - cx0) * ex / cex, oy + (y - cy0) * ey / cey)
                child.sx, child.sy = sx * ex / cex, sy * ey / cey
                self.shapes(el, part, slide_no, child, chrome, out, rels)
                continue
            if tag not in ('sp', 'pic', 'graphicFrame', 'cxnSp'):
                continue
            ph = X1(el, './/p:nvPr/p:ph')
            if chrome and ph is not None:
                continue   # layout placeholders are filled by the slide
            if ph is not None and ph.get('type') in ('dt', 'ftr') and not X(el, './/a:t'):
                continue
            xfrm = X1(el, './p:spPr/a:xfrm') if tag != 'graphicFrame' else X1(el, './p:xfrm')
            chain, inherited = (self._ph_chain(el, part) if tag == 'sp' else ([], None))
            if xfrm is None and inherited:
                for sp in inherited:
                    if sp is not None and X1(sp, './p:spPr/a:xfrm') is not None:
                        xfrm = X1(sp, './p:spPr/a:xfrm')
                        break
            if xfrm is None:
                continue
            off, ext = X1(xfrm, './a:off'), X1(xfrm, './a:ext')
            x, y = xf(int(off.get('x')), int(off.get('y')))
            w, h = int(ext.get('cx')) * sx, int(ext.get('cy')) * sy
            rot = int(xfrm.get('rot', '0')) / 60000
            flipH, flipV = xfrm.get('flipH') == '1', xfrm.get('flipV') == '1'
            name = X1(el, './/p:cNvPr/@name') or ''
            if tag == 'graphicFrame':
                gd = X1(el, './a:graphic/a:graphicData')
                uri = gd.get('uri', '')
                if uri.endswith('/chart'):
                    rid = X1(gd, './c:chart/@r:id')
                    cpart = part.related_part(rid)
                    chart = Chart(self, cpart.chart._chartSpace, w, h)
                    svg, label = chart.svg()
                    svg_m, _ = Chart(self, cpart.chart._chartSpace, w, h, mobile=True).svg()
                    cb = Box(x, y, w, h, svg, 'ch', chrome=chrome, text=label)
                    cb.data_table = getattr(chart, 'data_table', '')
                    cb.series_data = getattr(chart, 'series_data', [])
                    cb.html_m = svg_m
                    cb.cs, cb.deck = cpart.chart._chartSpace, self
                    out.append(cb)
                elif uri.endswith('/table'):
                    tb = Box(x, y, w, h, self.table(X1(gd, './a:tbl'), w, h, part, slide_no), 'tb', chrome=chrome)
                    tb.tbl, tb.part = X1(gd, './a:tbl'), part
                    out.append(tb)
                continue
            sppr = X1(el, './p:spPr')
            style = X1(el, './p:style')
            geom = X1(sppr, './a:prstGeom/@prst') or ('custom' if X1(sppr, './a:custGeom') is not None else 'rect')
            st = []
            if rot:
                st.append(f'transform:rotate({n(rot)}deg)')
            if tag == 'pic':
                blip = X1(el, './p:blipFill/a:blip')
                if blip is None or not blip.get(R + 'embed'):
                    continue
                src = X1(el, './p:blipFill/a:srcRect')
                crop = tuple(int(src.get(k, '0')) / 100000 for k in ('l', 't', 'r', 'b')) if src is not None else None
                if crop and not any(crop):
                    crop = None
                url = self.picture(blip, part, w, h, crop)
                alpha = X1(blip, './a:alphaModFix/@amt')
                descr = X1(el, './p:nvPicPr/p:cNvPr/@descr') or ''
                alt = '' if ('/' in descr or re.search(r'\.(png|jpe?g|svg)$', descr, re.I) or not descr) else descr
                if geom == 'ellipse':
                    st.append('border-radius:50%')
                elif geom == 'roundRect':
                    st.append(f'border-radius:calc(var(--u)*{n(min(w, h) * 0.16667)})')
                if alpha:
                    st.append(f'opacity:{n(int(alpha) / 100000)}')
                if flipH:
                    st.append('transform:scaleX(-1)')
                ln = self.theme.line(sppr, style)
                if ln:
                    st.append(f'border:max(1px,calc(var(--u)*{n(ln[0])})) {ln[2]} {ln[1]}')
                big = w * h > self.W * self.H * 0.08
                pb = Box(x, y, w, h, f'<img src="{esc(url)}" alt="{esc(alt)}" loading="lazy" decoding="async" style="{";".join(st)}">',
                               'pc' + (' pc-photo' if big and not alt else '') + (' pc-logo' if alt else ''), chrome=chrome,
                               deco=not alt, text=alt)
                pb.url, pb.alt, pb.geom = url, alt, geom
                out.append(pb)
                continue
            fill = self.theme.fill(sppr, style)
            ln = self.theme.line(sppr, style)
            txb = X1(el, './p:txBody')
            has_text = txb is not None and any((t.text or '').strip() for t in X(txb, './/a:t')) or (txb is not None and X(txb, './/a:fld[@type="slidenum"]'))
            if geom in ('line', 'straightConnector1', 'bentConnector2', 'bentConnector3') or tag == 'cxnSp':
                if not ln:
                    continue
                lw = f'max(1px,calc(var(--u)*{n(ln[0])}))'
                if h < 0.5 or w < 0.5:
                    side = 'top' if h < 0.5 else 'left'
                    out.append(Box(x, y, max(w, 0.01), max(h, 0.01), '', 'ln', chrome=chrome, deco=True))
                    out[-1].style = f'border-{side}:{lw} {ln[2]} {ln[1]}'
                else:
                    x1, x2 = (w, 0) if flipH else (0, w)
                    y1, y2 = (h, 0) if flipV else (0, h)
                    out.append(Box(x, y, w, h, f'<svg viewBox="0 0 {n(w)} {n(h)}" preserveAspectRatio="none" aria-hidden="true"><line x1="{n(x1)}" y1="{n(y1)}" x2="{n(x2)}" y2="{n(y2)}" stroke="{ln[1]}" stroke-width="{n(ln[0])}" vector-effect="non-scaling-stroke"/></svg>', 'ln', chrome=chrome, deco=True))
                continue
            if fill:
                st.append(f'background:{fill}')
            if ln:
                st.append(f'border:max(1px,calc(var(--u)*{n(ln[0])})) {ln[2]} {ln[1]}')
            if geom == 'ellipse':
                st.append('border-radius:50%')
            elif geom == 'roundRect':
                adj = X1(sppr, './a:prstGeom/a:avLst/a:gd[@name="adj"]/@fmla')
                k = int(adj.split()[-1]) / 100000 if adj else 0.16667
                st.append(f'border-radius:calc(var(--u)*{n(min(w, h) * k)})')
            elif geom == 'custom':
                if fill:
                    out.append(Box(x, y, w, h, self.custom(sppr, w, h, fill), 'cg', chrome=chrome, deco=True))
                continue
            inner, text, paras_meta = '', '', []
            if has_text:
                body = X1(txb, './a:bodyPr')
                lst = X1(txb, './a:lstStyle')
                bchain = [lst] + chain
                inh_body = [X1(sp, './p:txBody/a:bodyPr') for sp in (inherited or ()) if sp is not None]
                def battr(k, d):
                    for b in [body] + inh_body:
                        if b is not None and b.get(k) is not None:
                            return b.get(k)
                    return d
                ins = [int(battr(k, d)) / EMU_PT for k, d in (('tIns', 45720), ('rIns', 91440), ('bIns', 45720), ('lIns', 91440))]
                anchor = battr('anchor', 't')
                scale = 1.0
                fa = X1(body, './a:normAutofit') if body is not None else None
                if fa is not None and fa.get('fontScale'):
                    scale = int(fa.get('fontScale')) / 100000
                fref = self.theme.color(X1(style, './a:fontRef')) if style is not None else None
                is_title = ph is not None and ph.get('type') in ('title', 'ctrTitle')
                inner, text = self.text_html(txb, bchain, part.rels, slide_no, scale, fref, title=is_title)
                paras_meta = self.last_meta
                tst = [f'padding:{" ".join(f"calc(var(--u)*{n(v)})" if v else "0" for v in ins)}']
                if anchor in ('ctr', 'b'):
                    tst.append('display:flex;flex-direction:column;justify-content:' + ('center' if anchor == 'ctr' else 'flex-end'))
                if battr('wrap', 'square') == 'none':
                    tst.append('white-space:nowrap')
                cols = int(battr('numCol', '1'))
                if cols > 1:
                    gap = int(battr('spcCol', '0')) / EMU_PT
                    tst.append(f'column-count:{cols};column-gap:calc(var(--u)*{n(gap)});display:block')
                vert = battr('vert', 'horz')
                if vert in ('vert270', 'vert'):
                    tst.append('writing-mode:vertical-rl' + (';transform:rotate(180deg)' if vert == 'vert270' else ''))
                inner = f'<div class="tx" style="{";".join(tst)}">{inner}</div>'
            is_container = bool(fill or ln) and geom in ('rect', 'roundRect', 'ellipse') and min(w, h) > 8
            b = Box(x, y, w, h, inner, 'sh' + (' sh-t' if has_text else ''), container=is_container, chrome=chrome,
                    deco=not has_text, text=text)
            b.style = ';'.join(st)
            b.fill = fill
            b.line = ln
            b.paras = paras_meta if has_text else []
            b.ph = ph.get('type', 'body') if ph is not None else None
            b.geom = geom
            out.append(b)
        tag_links(mark)
        return out

    def custom(self, sppr, w, h, fill):
        paths = []
        for pth in X(sppr, './a:custGeom/a:pathLst/a:path'):
            pw, ph = int(pth.get('w', '1')) or 1, int(pth.get('h', '1')) or 1
            d = []
            for cmd in pth:
                k = etree.QName(cmd).localname
                pts = [(int(pt.get('x')) * w / pw, int(pt.get('y')) * h / ph) for pt in X(cmd, './a:pt')]
                if k == 'moveTo':
                    d.append('M%s %s' % tuple(n(v) for v in pts[0]))
                elif k == 'lnTo':
                    d.append('L%s %s' % tuple(n(v) for v in pts[0]))
                elif k == 'cubicBezTo':
                    d.append('C' + ' '.join(f'{n(a)} {n(b)}' for a, b in pts))
                elif k == 'quadBezTo':
                    d.append('Q' + ' '.join(f'{n(a)} {n(b)}' for a, b in pts))
                elif k == 'close':
                    d.append('Z')
            paths.append(' '.join(d))
        return f'<svg viewBox="0 0 {n(w)} {n(h)}" preserveAspectRatio="none" aria-hidden="true"><path d="{" ".join(paths)}" fill="{fill}"/></svg>'

    def table(self, tbl, w, h, part, slide_no):
        cols = [int(g.get('w')) / EMU_PT for g in X(tbl, './a:tblGrid/a:gridCol')]
        tw = sum(cols) or w
        rows = []
        for tr in X(tbl, './a:tr'):
            rh = int(tr.get('h', '0')) / EMU_PT
            cells = []
            for tc in X(tr, './a:tc'):
                if tc.get('hMerge') == '1' or tc.get('vMerge') == '1':
                    continue
                pr = X1(tc, './a:tcPr')
                st = []
                fill = self.theme.fill(pr) if pr is not None else None
                if fill:
                    st.append(f'background:{fill}')
                for side, k in (('left', 'lnL'), ('right', 'lnR'), ('top', 'lnT'), ('bottom', 'lnB')):
                    lnel = X1(pr, f'./a:{k}') if pr is not None else None
                    if lnel is not None and X1(lnel, './a:noFill') is None:
                        c = self.theme.color(X1(lnel, './a:solidFill'))
                        if c:
                            st.append(f'border-{side}:max(1px,calc(var(--u)*{n(int(lnel.get("w", 12700)) / EMU_PT)})) solid {c}')
                pads = [int(pr.get(k, d)) / EMU_PT if pr is not None else d / EMU_PT for k, d in (('marT', 45720), ('marR', 91440), ('marB', 45720), ('marL', 91440))]
                st.append('padding:' + ' '.join(f'calc(var(--u)*{n(v)})' for v in pads))
                anc = pr.get('anchor') if pr is not None else None
                st.append('vertical-align:' + {'ctr': 'middle', 'b': 'bottom'}.get(anc, 'top'))
                span = f' colspan="{tc.get("gridSpan")}"' if tc.get('gridSpan') else ''
                span += f' rowspan="{tc.get("rowSpan")}"' if tc.get('rowSpan') else ''
                inner, _ = self.text_html(X1(tc, './a:txBody'), [X1(tc, './a:txBody/a:lstStyle'), self.default_style], part.rels, slide_no)
                cells.append(f'<td{span} style="{";".join(st)}">{inner}</td>')
            rows.append(f'<tr style="height:calc(var(--u)*{n(rh)})">{"".join(cells)}</tr>')
        colg = ''.join(f'<col style="width:{n(c / tw * 100)}%">' for c in cols)
        return f'<table><colgroup>{colg}</colgroup><tbody>{"".join(rows)}</tbody></table>'

    # -- nesting -------------------------------------------------------------
    def tree(self, boxes):
        """Nests each element inside the smallest filled box under it that contains it."""
        roots = []
        placed = []
        for b in boxes:
            parent = None
            for c in reversed(placed):
                if not c.container or c is b:
                    continue
                tol = 1.0
                if (c.x - tol <= b.x and c.y - tol <= b.y and b.x + b.w <= c.x + c.w + tol and b.y + b.h <= c.y + c.h + tol
                        and (parent is None or c.w * c.h < parent.w * parent.h)):
                    parent = c
            (parent.children if parent else roots).append(b)
            placed.append(b)
        return roots


# ─── Charts ──────────────────────────────────────────────────────────────────
class Chart:
    def __init__(self, deck, cs, w, h, mobile=False, show_val=False, val_code=None):
        self.val_code = val_code  # number format for a value axis drawn here that the deck labels with its own text
        self.d, self.cs, self.w, self.h = deck, cs, w, h
        self.show_val = show_val  # draw the value axis even if the chart hides it (its labels sit beside the chart in the deck)
        self.t = deck.theme
        self.mobile = mobile      # a phone-width drawing: 340 units wide (about a CSS pixel each), larger text

    def tx(self, el, default_sz=9, default_color='#333'):
        """(size pt, color, weight, family) for a chart text properties element."""
        d = X1(el, './c:txPr//a:defRPr') if el is not None else None
        sz = int(d.get('sz')) / 100 if d is not None and d.get('sz') else default_sz
        if self.mobile:
            sz = max(11, sz * 1.45)
        col = self.t.color(X1(d, './a:solidFill')) if d is not None else None
        b = d is not None and d.get('b') == '1'
        face = X1(d, './a:latin/@typeface') if d is not None else None
        fam, wt = self.d.font(face or '+mn-lt', b)
        return sz, col or default_color, wt, fam

    def text(self, x, y, s, props, anchor='start', baseline='middle'):
        sz, col, wt, fam = props
        return (f'<text x="{n(x)}" y="{n(y)}" font-size="{n(sz)}" fill="{col}"{" font-weight=" + chr(34) + "700" + chr(34) if wt >= 600 else ""} '
                f'text-anchor="{anchor}" dominant-baseline="{baseline}">{esc(s)}</text>')

    def series(self, plot):
        cat_code_el = X(self.cs, './/c:dateAx/c:numFmt/@formatCode') + X(self.cs, './/c:catAx/c:numFmt/@formatCode')
        out = []
        for s in X(plot, './c:ser'):
            name = ''.join(X(s, './c:tx//c:v/text()'))
            cats = {int(pt.get('idx')): (X1(pt, './c:v/text()') or '') for pt in X(s, './c:cat//c:pt')}
            ccode = X1(s, './c:cat//c:formatCode/text()') or ''
            vals = {int(pt.get('idx')): float(X1(pt, './c:v/text()')) for pt in X(s, './c:val//c:pt') if X1(pt, './c:v/text()') not in (None, '')}
            count = int(X1(s, './c:val//c:ptCount/@val') or (max(vals) + 1 if vals else 0))
            code = next((c for c in cat_code_el if c and c != 'General'), ccode)
            labels = []
            for i in range(max(count, len(cats))):
                c = cats.get(i, '')
                if code and re.search(r'[ym]', code) and re.fullmatch(r'\d+(\.0)?', c or ''):
                    c = excel_date(float(c), code)
                labels.append(c)
            out.append(dict(el=s, name=name, cats=labels, vals=[vals.get(i) for i in range(len(labels))]))
        return out

    def scale(self, ax, lo, hi):
        mn, mx = X1(ax, './c:scaling/c:min/@val'), X1(ax, './c:scaling/c:max/@val')
        lo = float(mn) if mn is not None else min(0, lo)
        if mx is not None:
            hi = float(mx)
        else:
            span = hi - lo if hi > lo else abs(hi) or 1
            step = 10 ** __import__('math').floor(__import__('math').log10(span))
            for m in (1, 2, 2.5, 5, 10):
                if span / (step * m) <= 6:
                    step *= m
                    break
            hi = __import__('math').ceil(hi / step) * step if hi > 0 else 0
        if hi == lo:
            hi = lo + 1
        return lo, hi

    def svg(self):
        plot_area = X1(self.cs, './c:chart/c:plotArea')
        bar = X1(plot_area, './c:barChart')
        line = X1(plot_area, './c:lineChart')
        plot = bar if bar is not None else line
        W, H = self.w, self.h
        parts = []
        if plot is None:
            return f'<svg viewBox="0 0 {n(W)} {n(H)}" aria-hidden="true"></svg>', ''
        sers = self.series(plot)
        cats = sers[0]['cats'] if sers else []
        all_vals = [v for s in sers for v in s['vals'] if v is not None]
        cat_ax = X1(plot_area, './c:catAx') if X1(plot_area, './c:catAx') is not None else X1(plot_area, './c:dateAx')
        val_ax = X1(plot_area, './c:valAx')
        horiz = bar is not None and X1(bar, './c:barDir/@val') == 'bar'
        lo, hi = self.scale(val_ax, min(all_vals or [0]), max(all_vals or [1]))
        cat_rev = X1(cat_ax, './c:scaling/c:orientation/@val') == 'maxMin'
        val_rev = X1(val_ax, './c:scaling/c:orientation/@val') == 'maxMin'
        cat_del = X1(cat_ax, './c:delete/@val') == '1'
        val_del = X1(val_ax, './c:delete/@val') == '1'
        if self.mobile:
            nser = len(sers) if bar is not None else 1
            W = 340.0
            if horiz:
                H = 20 + max(1, len(cats)) * max(26, 16 * nser + 10)
            else:
                H = 230.0
            # Values that the deck shows as separate labels beside the chart are drawn on a value axis here.
            shows_values = any(X1(s['el'], './c:dLbls/c:showVal/@val') == '1' for s in sers) or X1(plot, './c:dLbls/c:showVal/@val') == '1'
            if val_del and not shows_values and self.show_val:
                val_del = False
        cat_tx = self.tx(cat_ax, 8)
        val_tx = self.tx(val_ax, 8)
        cat_lbl_pos = X1(cat_ax, './c:tickLblPos/@val') or 'nextTo'
        val_code = X1(val_ax, './c:numFmt/@formatCode') or 'General'
        skip = int(X1(cat_ax, './c:tickLblSkip/@val') or 0)
        if not skip and X1(cat_ax, './c:majorUnit/@val') and etree.QName(cat_ax).localname == 'dateAx':
            unit = float(X1(cat_ax, './c:majorUnit/@val'))
            skip = max(1, round(unit * (12 if X1(cat_ax, './c:majorTimeUnit/@val') == 'years' else 1)))
        skip = skip or 1
        # data labels
        def dl_props(s):
            d = X1(s['el'], './c:dLbls')
            if d is None:
                d = X1(plot, './c:dLbls')
            show = d is not None and X1(d, './c:showVal/@val') == '1'
            code = X1(d, './c:numFmt/@formatCode') if d is not None else None
            code = code or X1(s['el'], './c:val//c:formatCode/text()') or 'General'
            pos = X1(d, './c:dLblPos/@val') if d is not None else None
            deleted = {int(x.get('val')) for x in []}
            per = {}
            if d is not None:
                for dl in X(d, './c:dLbl'):
                    i = int(X1(dl, './c:idx/@val'))
                    if X1(dl, './c:delete/@val') == '1':
                        per[i] = None
                    else:
                        per[i] = X1(dl, './c:showVal/@val') == '1'
            return show, code, pos, self.tx(d, 8), per
        legend = X1(self.cs, './c:chart/c:legend')
        leg_pos = X1(legend, './c:legendPos/@val') if legend is not None else None
        leg_tx = self.tx(legend, 8)
        # plot area
        man = X1(plot_area, './c:layout/c:manualLayout')
        est = lambda s, sz: len(s) * sz * 0.5
        if man is not None and X1(man, './c:xMode/@val') != 'factor' and not self.mobile:
            px = float(X1(man, './c:x/@val') or 0) * W
            py = float(X1(man, './c:y/@val') or 0) * H
            pw = float(X1(man, './c:w/@val') or 1) * W
            ph = float(X1(man, './c:h/@val') or 1) * H
        else:
            top, bottom, left, right = 4, 4, 4, 4
            if legend is not None and leg_pos in ('t', 'b'):
                if leg_pos == 't':
                    top += leg_tx[0] * 1.8
                else:
                    bottom += leg_tx[0] * 1.8
            if legend is not None and leg_pos == 'r':
                right += max(est(s['name'], leg_tx[0]) for s in sers) + leg_tx[0] * 2
            any_dl = any(dl_props(s)[0] for s in sers)
            if horiz:
                if not cat_del and cat_lbl_pos != 'none':
                    left += min(W * 0.45, max([est(c, cat_tx[0]) for c in cats] or [0])) + 6
                if any_dl:
                    right += max(est(fmt(v, dl_props(sers[0])[1]), dl_props(sers[0])[3][0]) for v in all_vals) + 6 if all_vals else 0
                    pass
                if not val_del:
                    bottom += val_tx[0] * 1.6
            else:
                if not cat_del and cat_lbl_pos != 'none':
                    lines = 2 if bar is not None and max([est(c, cat_tx[0]) for c in cats] or [0]) > W / max(1, len(cats)) else 1
                    bottom += cat_tx[0] * 1.3 * lines + 4
                if any_dl:
                    top += dl_props(sers[0])[3][0] * 1.4
                if not val_del:
                    left += max(est(fmt(v, val_code), val_tx[0]) for v in (lo, hi)) + 6
            px, py, pw, ph = left, top, max(10, W - left - right), max(10, H - top - bottom)
        ncat = max(1, len(cats))
        # room for the labels of negative bars, which sit beyond the bar's end
        if horiz and bar is not None and all_vals and min(all_vals) < 0 and dl_props(sers[0])[0]:
            lw = max(est(fmt(v, dl_props(sers[0])[1]), dl_props(sers[0])[3][0]) for v in all_vals if v < 0) + 6
            if pw > lw * 2:
                lo_axis = lo
                lo = lo - (hi - lo) * lw / (pw - lw)
        # value -> position along the value axis
        def vpos(v):
            f = (v - lo) / (hi - lo)
            if val_rev:
                f = 1 - f
            return px + f * pw if horiz else py + ph - f * ph
        def cslot(i):
            j = (ncat - 1 - i) if (cat_rev != horiz) else i   # horizontal bars list the first category at the bottom unless reversed
            return j
        # gridlines and value axis
        grid = X1(val_ax, './c:majorGridlines')
        major = X1(val_ax, './c:majorUnit/@val')
        ticks = []
        if grid is not None or not val_del:
            step = float(major) if major else None
            if not step:
                span = hi - lo
                import math
                step = 10 ** math.floor(math.log10(span or 1))
                for m in (1, 2, 2.5, 5, 10):
                    if span / (step * m) <= 6:
                        step *= m
                        break
            import math
            # Excel starts the axis at its minimum; an axis drawn only for phones, or one widened for
            # negative labels, starts on a round value
            v = math.ceil(lo / step - 1e-9) * step if (self.show_val or 'lo_axis' in locals()) else lo
            while v <= hi + 1e-9 and len(ticks) < 40:
                ticks.append(v)
                v += step
        if grid is not None:
            gl = self.t.line(X1(grid, './c:spPr')) or (0.5, '#d9d9d9', 'solid')
            for v in ticks:
                p = vpos(v)
                if horiz:
                    parts.append(f'<line x1="{n(p)}" y1="{n(py)}" x2="{n(p)}" y2="{n(py + ph)}" stroke="{gl[1]}" stroke-width="{n(gl[0])}"/>')
                else:
                    parts.append(f'<line x1="{n(px)}" y1="{n(p)}" x2="{n(px + pw)}" y2="{n(p)}" stroke="{gl[1]}" stroke-width="{n(gl[0])}"/>')
        if self.val_code:
            val_code = self.val_code
        if not val_del and (X1(val_ax, './c:tickLblPos/@val') != 'none' or self.show_val):
            for v in ticks:
                p = vpos(v)
                s = fmt(v, val_code)
                if horiz:
                    parts.append(self.text(p, py + ph + val_tx[0] * 0.9, s, val_tx, 'middle'))
                else:
                    parts.append(self.text(px - 4, p, s, val_tx, 'end'))
        # bars
        slot = (pw if not horiz else ph) / ncat
        if bar is not None:
            gap = int(X1(bar, './c:gapWidth/@val') or 150) / 100
            ov = int(X1(bar, './c:overlap/@val') or 0) / 100
            k = len(sers)
            bw = slot / (k - ov * (k - 1) + gap)
            zero = vpos(max(lo, min(0, hi)))
            for si, s in enumerate(sers):
                sp = X1(s['el'], './c:spPr')
                base = self.t.fill(sp) or '#888'
                dpt = {int(X1(d, './c:idx/@val')): self.t.fill(X1(d, './c:spPr')) for d in X(s['el'], './c:dPt')}
                show, code, pos, dtx, per = dl_props(s)
                for i, v in enumerate(s['vals']):
                    if v is None:
                        continue
                    j = cslot(i)
                    start = (px if not horiz else py) + j * slot + slot * gap / 2 / (1 + 0) * (1) / (1) * 0 + (slot - bw * (k - ov * (k - 1))) / 2 + si * bw * (1 - ov)
                    p = vpos(v)
                    color = dpt.get(i) or base
                    if horiz:
                        x0, x1 = sorted((zero, p))
                        parts.append(f'<rect x="{n(x0)}" y="{n(start)}" width="{n(max(x1 - x0, 0.01))}" height="{n(bw)}" fill="{color}"/>')
                    else:
                        y0, y1 = sorted((zero, p))
                        parts.append(f'<rect x="{n(start)}" y="{n(y0)}" width="{n(bw)}" height="{n(max(y1 - y0, 0.01))}" fill="{color}"/>')
                    want = per.get(i, show) if i in per else show
                    if want:
                        lab = fmt(v, code)
                        if horiz:
                            neg = v < 0
                            if pos in ('inEnd',):
                                parts.append(self.text(p + (4 if neg else -4), start + bw / 2, lab, dtx, 'start' if neg else 'end'))
                            else:
                                parts.append(self.text(p + (-3 if neg else 3), start + bw / 2, lab, dtx, 'end' if neg else 'start'))
                        else:
                            neg = v < 0
                            parts.append(self.text(start + bw / 2, p + (dtx[0] * 0.8 if neg else -dtx[0] * 0.7), lab, dtx, 'middle'))
        # lines
        if line is not None:
            for s in sers:
                sp = X1(s['el'], './c:spPr')
                ln = self.t.line(sp) or (1.5, '#888', 'solid')
                pts = []
                for i, v in enumerate(s['vals']):
                    if v is None:
                        continue
                    j = cslot(i) if horiz else (ncat - 1 - i if cat_rev else i)
                    cx = px + (j + 0.5) * slot
                    pts.append(f'{n(cx)},{n(vpos(v))}')
                dash = ' stroke-dasharray="4 3"' if ln[2] == 'dashed' else ''
                parts.append(f'<polyline points="{" ".join(pts)}" fill="none" stroke="{ln[1]}" stroke-width="{n(ln[0])}" stroke-linejoin="round"{dash}/>')
                mk = X1(s['el'], './c:marker')
                sym = X1(mk, './c:symbol/@val') if mk is not None else 'none'
                if sym and sym != 'none':
                    msz = int(X1(mk, './c:size/@val') or 5) * 0.75
                    mfill = self.t.fill(X1(mk, './c:spPr')) or ln[1]
                    mln = self.t.line(X1(mk, './c:spPr'))
                    stroke = f' stroke="{mln[1]}" stroke-width="{n(mln[0])}"' if mln else ''
                    r = msz / 2
                    for pt in pts:
                        cx, cy = (float(v) for v in pt.split(','))
                        if sym == 'circle':
                            parts.append(f'<circle cx="{n(cx)}" cy="{n(cy)}" r="{n(r)}" fill="{mfill}"{stroke}/>')
                        elif sym == 'diamond':
                            parts.append(f'<path d="M{n(cx)} {n(cy - r * 1.2)}L{n(cx + r * 1.2)} {n(cy)}L{n(cx)} {n(cy + r * 1.2)}L{n(cx - r * 1.2)} {n(cy)}Z" fill="{mfill}"{stroke}/>')
                        elif sym == 'triangle':
                            parts.append(f'<path d="M{n(cx)} {n(cy - r)}L{n(cx + r)} {n(cy + r)}L{n(cx - r)} {n(cy + r)}Z" fill="{mfill}"{stroke}/>')
                        else:
                            parts.append(f'<rect x="{n(cx - r)}" y="{n(cy - r)}" width="{n(2 * r)}" height="{n(2 * r)}" fill="{mfill}"{stroke}/>')
                show, code, pos, dtx, per = dl_props(s)
                for i, v in enumerate(s['vals']):
                    want = per.get(i, show) if i in per else show
                    if v is None or not want:
                        continue
                    j = ncat - 1 - i if cat_rev else i
                    cx = px + (j + 0.5) * slot
                    parts.append(self.text(cx, vpos(v) - dtx[0] * 0.8, fmt(v, code), dtx, 'middle'))
        # category axis line and labels
        ax_ln = self.t.line(X1(cat_ax, './c:spPr'))
        if ax_ln and not cat_del:
            if horiz:
                zx = vpos(max(lo, min(0, hi)))
                parts.append(f'<line x1="{n(zx)}" y1="{n(py)}" x2="{n(zx)}" y2="{n(py + ph)}" stroke="{ax_ln[1]}" stroke-width="{n(ax_ln[0])}"/>')
            else:
                zy = vpos(max(lo, min(0, hi)))
                parts.append(f'<line x1="{n(px)}" y1="{n(zy)}" x2="{n(px + pw)}" y2="{n(zy)}" stroke="{ax_ln[1]}" stroke-width="{n(ax_ln[0])}"/>')
        if not cat_del and cat_lbl_pos != 'none':
            if not horiz and cats:
                widest = max(len(c) for c in cats) * cat_tx[0] * 0.55 + 4
                while skip < len(cats) and slot * skip < widest:
                    skip += 1
            for i, c in enumerate(cats):
                if i % skip:
                    continue
                j = cslot(i) if bar is not None else (ncat - 1 - i if cat_rev else i)
                mid = (px if not horiz else py) + (j + 0.5) * slot
                if horiz:
                    parts.append(self.wrap_text(px - 6, mid, c, cat_tx, 'end', px - 6))
                else:
                    y = (py + ph) + cat_tx[0] * 0.9 + 2
                    parts.append(self.wrap_text(mid, y, c, cat_tx, 'middle', slot * max(1, skip) * 0.98, top=True))
        # legend
        if legend is not None and X1(legend, './c:delete/@val') != '1':
            items = []
            for s in sers:
                sp = X1(s['el'], './c:spPr')
                col = (self.t.fill(sp) if bar is not None else (self.t.line(sp) or (0, '#888'))[1]) or '#888'
                items.append((s['name'], col))
            sz = leg_tx[0]
            if leg_pos == 'r':
                lx, ly = px + pw + sz, py + sz
                for name, col in items:
                    parts.append(f'<rect x="{n(lx)}" y="{n(ly - sz * 0.35)}" width="{n(sz * 0.7)}" height="{n(sz * 0.7)}" fill="{col}"/>')
                    parts.append(self.text(lx + sz, ly, name, leg_tx))
                    ly += sz * 1.6
            else:
                total = sum(est(nm, sz) + sz * 2.2 for nm, _ in items)
                lx = max(2, (W - total) / 2)
                ly = (sz * 0.9) if leg_pos == 't' else H - sz * 0.9
                for name, col in items:
                    parts.append(f'<rect x="{n(lx)}" y="{n(ly - sz * 0.35)}" width="{n(sz * 0.7)}" height="{n(sz * 0.7)}" fill="{col}"/>')
                    parts.append(self.text(lx + sz, ly, name, leg_tx))
                    lx += est(name, sz) + sz * 2.2
        # accessible summary
        label = '; '.join(f'{s["name"] + ": " if s["name"] else ""}' + ', '.join(
            f'{c} {fmt(v, dl_props(s)[1] if dl_props(s)[0] else val_code)}' for c, v in zip(s['cats'], s['vals']) if v is not None) for s in sers)
        svg = (f'<svg viewBox="0 0 {n(W)} {n(H)}" role="img" aria-label="{esc(label[:900])}" preserveAspectRatio="xMidYMid meet">'
               + ''.join(parts) + '</svg>')
        # the chart's values as a table, for readers, search engines and AI tools
        codes = [dl_props(s)[1] if dl_props(s)[0] else (X1(s['el'], './c:val//c:formatCode/text()') if val_code == 'General' else val_code) for s in sers]
        head = '<tr><th scope="col"></th>' + ''.join(f'<th scope="col">{esc(s["name"])}</th>' for s in sers) + '</tr>'
        rows = ''.join('<tr><th scope="row">' + esc(c) + '</th>' + ''.join(
            f'<td>{esc(fmt(s["vals"][i], codes[k]) if i < len(s["vals"]) else "")}</td>' for k, s in enumerate(sers)) + '</tr>'
            for i, c in enumerate(cats))
        self.data_table = f'<table><thead>{head}</thead><tbody>{rows}</tbody></table>'
        self.series_data = [(s['name'], s['cats'], [fmt(v, codes[k]) if v is not None else '' for v in s['vals']]) for k, s in enumerate(sers)]
        return svg, ''

    def wrap_text(self, x, y, s, props, anchor, maxw, top=False):
        sz = props[0]
        words, lines, cur = s.split(), [], ''
        for wd in words:
            t = (cur + ' ' + wd).strip()
            if cur and len(t) * sz * 0.5 > maxw:
                lines.append(cur)
                cur = wd
            else:
                cur = t
        if cur:
            lines.append(cur)
        if len(lines) <= 1:
            return self.text(x, y, s, props, anchor)
        lh = sz * 1.15
        y0 = y if top else y - lh * (len(lines) - 1) / 2
        return ''.join(self.text(x, y0 + i * lh, ln, props, anchor) for i, ln in enumerate(lines))
